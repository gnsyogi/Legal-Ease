from docx import Document
from fpdf import FPDF

def format_docx(text, filename="output.docx"):
    doc = Document()
    doc.add_heading("Legal Document", 0)
    doc.add_paragraph(text)
    doc.save(filename)
    return filename

def format_pdf(text, filename="output.pdf"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for line in text.split("\n"):
        pdf.multi_cell(0, 8, line)

    pdf.output(filename)
    return filename
